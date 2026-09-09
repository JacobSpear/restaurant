"""
Bi-weekly report generation and Excel export.
"""

from __future__ import annotations

import datetime
from datetime import timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, NamedStyle, PatternFill

from . import config as cfg
from .summaries import report_table
from .utils import one_year_prior, two_weeks_prior

pd.options.mode.chained_assignment = None


# ---------------------------------------------------------------------------
# Two-week report helpers
# ---------------------------------------------------------------------------

def two_week_report(
    start_day: str,
    orders: pd.DataFrame,
    sales: pd.DataFrame,
    area: int = 0,
) -> pd.DataFrame | None:
    """Build a two-week summary starting on *start_day* (must be a Monday).

    Returns ``None`` if *start_day* is not a Monday.
    """
    sdy1 = datetime.datetime.strptime(start_day, "%Y-%m-%d").date()
    if sdy1.weekday() != 0:
        return None

    edy1 = sdy1 + timedelta(days=6)
    sdy2 = sdy1 + timedelta(days=7)
    edy2 = sdy1 + timedelta(days=13)

    w1 = report_table(
        sdy1.strftime("%Y-%m-%d"), edy1.strftime("%Y-%m-%d"),
        orders, sales, area,
    ).rename(columns={"Total": "Week 1 Total"})
    w2 = report_table(
        sdy2.strftime("%Y-%m-%d"), edy2.strftime("%Y-%m-%d"),
        orders, sales, area,
    ).rename(columns={"Total": "Week 2 Total"})
    both = report_table(
        sdy1.strftime("%Y-%m-%d"), edy2.strftime("%Y-%m-%d"),
        orders, sales, area,
    )[["index", "Total"]]

    return (
        pd.merge(pd.merge(w1, w2, on="index", how="outer"), both, on="index", how="outer")
    )


def _pct_change_str(current_val: float, base_val: float) -> str:
    """Format ``base -> current (+X%)`` change string."""
    base = round(base_val, 2)
    comp = round(current_val, 2)
    if base > 0:
        pct = 100 * (comp - base) / base
        sign = "+" if pct >= 0 else "-"
        return f"${base} -> ${comp} ({sign}{abs(round(pct, 1))}%)"
    return f"{base}->{comp}"


def two_week_change(
    current_day: str,
    comp_day: str,
    row_labels: list[str],
    orders: pd.DataFrame,
    sales: pd.DataFrame,
    area: int = 0,
) -> pd.DataFrame:
    """Compare two-week totals between *current_day* and *comp_day*."""
    cur = two_week_report(current_day, orders, sales, area)[
        ["index", "Week 1 Total", "Week 2 Total", "Total"]
    ]
    comp = two_week_report(comp_day, orders, sales, area)[
        ["index", "Week 1 Total", "Week 2 Total", "Total"]
    ]
    cur2 = cur.loc[cur["index"].isin(row_labels)]
    comp2 = comp.loc[comp["index"].isin(row_labels)]

    out: dict[str, list] = {"index": [], "Week 1 Total": [], "Week 2 Total": [], "Total": []}
    for row in row_labels:
        out["index"].append(row)
        for col in ("Week 1 Total", "Week 2 Total", "Total"):
            c_val = float(cur2.loc[cur2["index"] == row, col])
            b_val = float(comp2.loc[comp2["index"] == row, col])
            out[col].append(_pct_change_str(c_val, b_val))
    return pd.DataFrame(out)


_DAY_COLUMNS = [
    "index",
    "Week 1: Monday", "Week 1: Tuesday", "Week 1: Wednesday",
    "Week 1: Thursday", "Week 1: Friday", "Week 1: Saturday",
    "Week 1: Sunday", "Week 1 Total",
    "Week 2: Monday", "Week 2: Tuesday", "Week 2: Wednesday",
    "Week 2: Thursday", "Week 2: Friday", "Week 2: Saturday",
    "Week 2: Sunday", "Week 2 Total",
    "Total",
]


def daily_change(
    current_day: str,
    comp_day: str,
    row_labels: list[str],
    orders: pd.DataFrame,
    sales: pd.DataFrame,
    area: int = 0,
) -> pd.DataFrame:
    """Compare daily columns between two bi-weekly periods."""
    cur = two_week_report(current_day, orders, sales, area)
    comp = two_week_report(comp_day, orders, sales, area)
    cur.columns = _DAY_COLUMNS
    comp.columns = _DAY_COLUMNS
    cur2 = cur.loc[cur["index"].isin(row_labels)]
    comp2 = comp.loc[comp["index"].isin(row_labels)]

    out: dict[str, list] = {c: [] for c in _DAY_COLUMNS}
    for row in row_labels:
        out["index"].append(row)
        for col in set(_DAY_COLUMNS) - {"index"}:
            c_val = float(cur2.loc[cur2["index"] == row, col])
            b_val = float(comp2.loc[comp2["index"] == row, col])
            out[col].append(_pct_change_str(c_val, b_val))
    return pd.DataFrame(out)


def _add_pct_column(df: pd.DataFrame, src_col: str, new_col: str) -> None:
    """Insert a '% of Non-Service Total' column after *src_col*."""
    vals = dict(zip(df["index"], df[src_col]))
    total = vals.get("Total", 0)
    service = vals.get("Total Service", 0)
    denom = total - service
    col_data = [
        "" if k in ("Total", "Total Service", "Tock Sales-Discounts", "Toast Net")
        else (f"{round(100 * v / denom, 2)}%" if denom != 0 else "")
        for k, v in vals.items()
    ]
    pos = df.columns.get_loc(src_col) + 1
    df.insert(pos, new_col, col_data)


def _add_pct_columns(df: pd.DataFrame) -> None:
    """Add percentage-of-total columns to a bi-weekly summary."""
    _add_pct_column(df, "Week 1 Total", "Week 1 % of Non-Service Total")
    _add_pct_column(df, "Week 2 Total", "Week 2 % of Non-Service Total")
    _add_pct_column(df, "Total", "% of Non-Service Total")


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def export_to_excel(
    filepath: str | Path,
    dataframes: list[pd.DataFrame],
    sheet_titles: list[str],
) -> None:
    """Write *dataframes* to an Excel workbook with currency/percentage formatting."""
    filepath = Path(filepath)
    filepath.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        for df, title in zip(dataframes, sheet_titles):
            df.to_excel(writer, sheet_name=title, index=False)

    wb = load_workbook(filepath)
    pink = PatternFill(start_color="E49EDD", end_color="E49EDD", fill_type="solid")
    green = PatternFill(start_color="B5E6A2", end_color="B5E6A2", fill_type="solid")

    if "currency" not in wb.named_styles:
        ns = NamedStyle(name="currency")
        ns.number_format = "$#,##0.00"
        wb.add_named_style(ns)
    if "percent" not in wb.named_styles:
        ns = NamedStyle(name="percent")
        ns.number_format = "##0.00%"
        wb.add_named_style(ns)

    for df, title in zip(dataframes, sheet_titles):
        ws = wb[title]
        ws["A1"] = title
        ws["A1"].font = Font(bold=True)

        # Identify highlight columns
        col_fills: dict[int, PatternFill] = {}
        for idx, cell in enumerate(ws[1], start=1):
            val = cell.value or ""
            if "Week" in val and "Weeks" not in val:
                col_fills[idx] = pink
            elif "Total" in val:
                col_fills[idx] = green

        # Format data cells
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.style = "currency"
                elif isinstance(cell.value, str):
                    if "$" in cell.value:
                        try:
                            cell.value = float(
                                cell.value.replace("$", "").replace(",", "")
                            )
                            cell.style = "currency"
                        except ValueError:
                            pass
                    elif "%" in cell.value:
                        try:
                            cell.value = (
                                float(cell.value.replace("%", "").replace(",", ""))
                                / 100
                            )
                            cell.style = "percent"
                        except ValueError:
                            pass

        # Column fills
        for col_idx, fill in col_fills.items():
            for row in ws.iter_rows(
                min_row=1, max_row=ws.max_row, min_col=col_idx, max_col=col_idx
            ):
                for cell in row:
                    cell.fill = fill

        # Fix escaped dollar signs
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and "\$" in cell.value:
                    cell.value = cell.value.replace("\$", "$")

        # Auto-width
        for col in ws.columns:
            letter = col[0].column_letter
            width = max(len(str(c.value or "")) for c in col) + 2
            ws.column_dimensions[letter].width = width

    wb.save(filepath)


# ---------------------------------------------------------------------------
# Full bi-weekly report
# ---------------------------------------------------------------------------

CHANGE_ROWS = [
    "Tock Sales-Discounts", "Toast Net", "Total Bev", "Wine+Pairings",
    "Toast Beer", "Toast Liquor", "Total Food", "Total Service", "Total",
]


def generate_biweekly_report(
    start_monday: str,
    orders: pd.DataFrame,
    sales: pd.DataFrame,
    output_dir: str | Path = ".",
) -> None:
    """Generate the full bi-weekly report suite (CSVs + Excel)."""
    output_dir = Path(output_dir)

    # Summaries by area
    summaries = {}
    area_labels = {0: "All", 1: "DR", 2: "BAR", 3: "BB"}
    for area, label in area_labels.items():
        df = two_week_report(start_monday, orders, sales, area)
        df["idx"] = df["index"].map(cfg.REPORT_ROW_ORDER)
        df = (
            df.loc[df["index"].isin(cfg.REPORT_ROWS)]
            .round(2)
            .sort_values("idx")
            .drop(columns="idx")
        )
        _add_pct_columns(df)
        summaries[label] = df

    # Change reports
    chg_2wk = two_week_change(
        start_monday, two_weeks_prior(start_monday), CHANGE_ROWS, orders, sales,
    ).round(2)
    chg_yr = two_week_change(
        start_monday, one_year_prior(start_monday), CHANGE_ROWS, orders, sales,
    ).round(2)
    dchg_2wk = daily_change(
        start_monday, two_weeks_prior(start_monday), CHANGE_ROWS, orders, sales,
    ).round(2)
    dchg_yr = daily_change(
        start_monday, one_year_prior(start_monday), CHANGE_ROWS, orders, sales,
    ).round(2)

    # Compute report-date label
    end_date = (
        datetime.datetime.strptime(start_monday, "%Y-%m-%d") + timedelta(days=13)
    ).strftime("%d%b%Y").upper()
    release_date = (
        datetime.datetime.strptime(end_date, "%d%b%Y") + timedelta(days=5)
    ).strftime("%Y-%m-%d")
    report_dir = output_dir / release_date
    report_dir.mkdir(parents=True, exist_ok=True)

    # CSV exports
    for label, df in summaries.items():
        df.to_csv(report_dir / f"BiWeeklySales_{label}.csv", index=True,
                  float_format="${:,.2f}".format)
    chg_2wk.to_csv(report_dir / "BiWeeklyChangeP_All.csv", index=True)
    chg_yr.to_csv(report_dir / "BiWeeklyChangeY_All.csv", index=True)

    # Excel
    all_dfs = [
        summaries["All"], summaries["DR"], summaries["BAR"], summaries["BB"],
        chg_2wk, chg_yr, dchg_2wk, dchg_yr,
    ]
    sheet_names = [
        "Whole Restaurant", "Dining Room", "Bar", "Bayani Bar",
        "Change - 2 Weeks", "Change - Year",
        "Daily Change - 2 Weeks", "Daily Change - Year",
    ]
    export_to_excel(
        report_dir / f"BiWeeklySummary_{end_date}.xlsx",
        all_dfs, sheet_names,
    )
    print(f"Bi-weekly report written to {report_dir}/")
